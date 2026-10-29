"""Export and verify saved grouped reports without rerunning scientific analysis.

The exported workbook uses the optional openpyxl dependency. Verification reads
every parent/candidate cell and the native filter/freeze metadata without
repairing or otherwise changing the workbook under inspection.
"""
from __future__ import annotations

import argparse
from contextlib import closing
import csv
from hashlib import sha256
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile


def manifest_path(batch):
    selected = batch / "selected_report_manifest.json"
    if selected.is_file():
        return selected
    ordinary = batch / "batch_manifest.json"
    return ordinary if ordinary.is_file() else batch / "combined_manifest.json"


def combined_csv_path(batch, manifest):
    """Honor a user rename only when exactly one file retains the saved hash."""
    from fv_energy_analysis.paths import native_path
    recorded = native_path(manifest["csv"])
    if recorded.is_file():
        assert sha256(recorded.read_bytes()).hexdigest() == manifest["csv_sha256"], "Combined CSV hash differs"
        return recorded
    matches = [path for path in batch.glob("*.csv")
               if sha256(path.read_bytes()).hexdigest() == manifest["csv_sha256"]]
    assert len(matches) == 1, "Expected exactly one combined CSV matching its saved hash"
    return matches[0]


def combined_failure_details(batch, rows):
    """Resolve saved source diagnostics without recalculating ligand chemistry."""
    from fv_energy_analysis.paths import native_path
    with (batch / "candidate_sources.csv").open(encoding="utf-8-sig", newline="") as stream:
        origins = list(csv.DictReader(stream))
    assert len(origins) == len(rows), "Candidate provenance count differs from combined rows"
    assert len({r["candidate"] for r in origins}) == len(rows), "Duplicate candidate provenance"
    assert {r["candidate"] for r in origins} == {r["candidate"] for r in rows}, "Candidate provenance inventory differs"
    by_name = {r["candidate"]: r for r in origins}
    source_cases = {}
    failures = []
    for row in rows:
        if row["status"] != "failed":
            continue
        source = native_path(by_name[row["candidate"]]["source"])
        if source not in source_cases:
            saved = source / "preparation_failures.json"
            cases = json.loads(saved.read_text(encoding="utf-8")) if saved.is_file() else []
            source_cases[source] = {case["candidate"]: case for case in cases}
        case = source_cases[source].get(row["candidate"], {})
        diagnostic = native_path(case["diagnostic_folder"]) / "ligand_identity_diagnostic.json" if case.get("diagnostic_folder") else None
        detail = json.loads(diagnostic.read_text(encoding="utf-8")) if diagnostic and diagnostic.is_file() else {}
        error = row.get("error") or "Analysis did not complete"
        if detail.get("differences") == ["stereochemistry"] and detail.get("stereocentre_differences"):
            error = "Stereochemistry mismatch: " + "; ".join(
                f"{c['atom_name']}: expected {c['expected_CIP']}, observed {c['observed_CIP']}"
                for c in detail["stereocentre_differences"])
        failures.append({"candidate": row["candidate"], "error": error,
                         "expected": detail.get("expected", {}).get("inchikey"),
                         "observed": detail.get("observed", {}).get("inchikey"),
                         "diagnostic": str(diagnostic) if detail else str(source), "source": str(source)})
    return failures


def typed(value, column):
    if value is None or value in ("", "NA"):
        return None
    kind, key = column["kind"], column["key"]
    if kind == "text":
        return str(value)
    if kind == "boolean":
        if isinstance(value, bool):
            return value
        if str(value).lower() in ("true", "false"):
            return str(value).lower() == "true"
        raise ValueError(f"Invalid boolean: {key}={value!r}")
    if isinstance(value, bool) or kind not in ("integer", "number", "energy", "delta_energy"):
        raise ValueError(f"Invalid numeric kind/value: {key}={value!r}")
    number = float(value)
    if not math.isfinite(number) or (kind == "integer" and not number.is_integer()):
        raise ValueError(f"Invalid number: {key}={value!r}")
    return int(number) if kind == "integer" else number


def prepare(batch):
    batch = Path(batch).resolve()
    if (batch / "selected_report_manifest.json").is_file():
        from .report_options import prepare_selected_report
        return prepare_selected_report(batch)
    from fv_energy_analysis.confidence import confidence_columns
    from fv_energy_analysis.fingerprint_reports import _grouped_columns
    from fv_energy_analysis.model_scores import score_columns

    payload = json.loads((batch / "batch_summary_grouped_fingerprints.json").read_text(encoding="utf-8-sig"))
    selected_manifest = manifest_path(batch)
    manifest = json.loads(selected_manifest.read_text(encoding="utf-8-sig"))
    combined = selected_manifest.name == "combined_manifest.json"
    columns, rows = payload["columns"], payload["rows"]
    present_sections = {column["section"] for column in columns}
    legacy = payload.get("scoring", {}).get("legacy_pre_talaris")
    if legacy is None:
        legacy = manifest.get("scientific_protocol", {}).get("legacy_pre_talaris")
    if legacy is None:
        legacy = payload.get("scoring_profile") == "rosettaligand_legacy"
    expected_columns = _grouped_columns("ligand" if legacy else "ref2015")
    if "HCDR3 pLDDT" in present_sections:
        expected_columns += confidence_columns()
    if present_sections & {"Boltz-2", "GNINA"}:
        expected_columns += score_columns()
    if combined:
        expected_columns += [
            {"key": "status", "label": "Analysis status", "section": "Analysis status", "kind": "text"},
            {"key": "error", "label": "Failure reason", "section": "Analysis status", "kind": "text"},
        ]
    if payload.get("e3_discrimination_protocol"):
        from fv_energy_analysis.e3_reports import e3_columns
        insert = next((i + 1 for i in range(len(expected_columns)-1, -1, -1)
                       if expected_columns[i]["section"] == "E2 architecture"), len(expected_columns))
        expected_columns[insert:insert] = e3_columns()
    metadata = lambda cols: [tuple(c[key] for key in ("key", "label", "section", "kind")) for c in cols]
    assert metadata(columns) == metadata(expected_columns), "Report schema differs from the current package"
    assert len({c["key"] for c in columns}) == len(columns), "Duplicate column keys"
    expected_sections = list(dict.fromkeys(c["section"] for c in expected_columns))
    sections = list(dict.fromkeys(c["section"] for c in columns))
    assert sections == expected_sections, f"Unexpected section order: {sections}"
    if payload.get("e3_discrimination_protocol", {}).get("source_combined_batch"):
        from fv_energy_analysis.paths import native_path
        protocol = payload["e3_discrimination_protocol"]
        original_path = native_path(protocol["source_combined_batch"]) / "batch_summary_grouped_fingerprints.json"
        assert sha256(original_path.read_bytes()).hexdigest() == protocol["source_payload_sha256"], "E3 source data changed"
        original = json.loads(original_path.read_text(encoding="utf-8-sig"))
        original_keys = [c["key"] for c in original["columns"]]
        assert [r["candidate"] for r in rows] == [r["candidate"] for r in original["rows"]], "E3 candidate order changed"
        for old, new in zip([original["parent"], *original["rows"]], [payload["parent"], *rows]):
            assert all(old.get(key) == new.get(key) for key in original_keys), "E3 export changed an original metric"
    names = [r["candidate"] for r in rows]
    assert len(names) == len(set(names)), "Duplicate candidate identifiers"
    if payload.get("architecture_protocol"):
        from fv_energy_analysis.paths import native_path
        protocol = payload["architecture_protocol"]
        assert protocol["replicates"] == 1 and protocol["arpeggio_used"] is False
        assert protocol["interaction_source"] == "Discovery Studio"
        original = json.loads((native_path(protocol["source_batch"]) /
                               "batch_summary_grouped_fingerprints.json").read_text(encoding="utf-8"))
        unchanged = [c for c in original["columns"] if c["section"] != "E2 architecture"]
        assert metadata(unchanged) == metadata([c for c in columns if c["section"] != "E2 architecture"])
        assert [r["candidate"] for r in original["rows"]] == names
        for before, after in zip([original["parent"], *original["rows"]], [payload["parent"], *rows]):
            assert all(before.get(c["key"]) == after.get(c["key"]) for c in unchanged), "Unrelated metrics changed"
            expected = sum(after[k] for k in ("e2_o3_hbond", "e2_o17_hbond",
                                             "e2_c18_hydrophobic_contact", "e2_geometric_sandwich"))
            assert after["e2_conserved_architecture_score_0_to_4"] == expected
            assert after["e2_strict_screen_pass"] is (expected == 4)
            assert after["e2_c18_hydrophobic_contact"] is (after["e2_c18_hydrophobic_partner_count"] >= 1)
            steric = after["e2_steric_opposing_residue_pair_count"] >= 1
            assert after["e2_geometric_sandwich"] is steric
    for row in [payload.get("parent", {}), *rows]:
        for column in columns:
            typed(row.get(column["key"]), column)
    if combined:
        csv_path = combined_csv_path(batch, manifest)
        assert sha256(csv_path.read_bytes()).hexdigest() == manifest["csv_sha256"], "Combined CSV hash differs"
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            records = list(reader)
            assert reader.fieldnames == [c["key"] for c in columns], "Combined CSV/JSON column order differs"
        assert len(records) == len(rows) == int(manifest["candidate_count"]), "Combined row count differs"
        for index, (actual, saved) in enumerate(zip(records, rows), 1):
            assert all(typed(actual.get(c["key"]), c) == typed(saved.get(c["key"]), c) for c in columns), f"Combined CSV/JSON row differs: {index}"
        completed = [r for r in rows if r["status"] == "complete"]
        rejected = [r for r in rows if r["status"] == "failed"]
        assert len(completed) == int(manifest["completed_count"]), "Combined completed count differs"
        assert len(rejected) == int(manifest["failure_count"]), "Combined rejection count differs"
        assert len(completed) + len(rejected) == len(rows), "Unknown combined status"
        assert [r["rank"] for r in completed] == list(range(1, len(completed) + 1)), "Combined global ranks differ"
        assert all(r.get("rank") is None and r.get("rosetta_interface_REU") is None and r.get("rosetta_interface_delta_REU") is None and r.get("error") for r in rejected), "Rejected candidate contains scores or lacks its reason"
        return payload, manifest, combined_failure_details(batch, rows)
    failures = []
    diagnostics = {}
    diagnostic_csv = batch / "ligand_validation_diagnostics.csv"
    if diagnostic_csv.is_file():
        with diagnostic_csv.open(encoding="utf-8-sig", newline="") as stream:
            diagnostics = {row["candidate"]: row for row in csv.DictReader(stream)}
    if int(manifest.get("failure_count", 0)):
        with (batch / "batch_summary.csv").open(encoding="utf-8-sig", newline="") as stream:
            records = list(csv.DictReader(stream))
        failures = [{"candidate": Path(r["candidate"].replace("\\", "/")).name,
                     "error": r.get("error") or "Analysis did not complete",
                     "source": r.get("input_path") or r.get("candidate_path") or r["candidate"]}
                    for r in records if r.get("status") != "complete"]
    assert len(rows) == int(manifest["completed_count"]), "Completed count differs from report rows"
    for failure in failures:
        detail = diagnostics.get(failure["candidate"], {})
        if detail.get("differences") == "stereochemistry":
            centres = json.loads(detail["stereocentre_differences"])
            failure["error"] = "Stereochemistry mismatch: " + "; ".join(
                f"{c['atom_name']}: expected {c['expected_CIP']}, observed {c['observed_CIP']}" for c in centres)
        failure["expected"] = detail.get("expected_inchikey") or None
        failure["observed"] = detail.get("observed_inchikey") or None
        failure["diagnostic"] = detail.get("diagnostic_path") or failure["source"]
    assert len(failures) == int(manifest.get("failure_count", 0)), "Failure count differs from detailed records"
    assert len(rows) + len(failures) == int(manifest["candidate_count"]), "Candidates were silently omitted"
    assert not set(names).intersection(f["candidate"] for f in failures), "Candidate both complete and failed"
    return payload, manifest, failures


def workbook_path(batch, manifest, explicit=None):
    """Resolve an existing workbook; a user rename requires an unambiguous file."""
    from .paths import native_path

    if explicit is not None:
        result = native_path(explicit).resolve()
        if not result.is_file():
            raise FileNotFoundError(result)
        return result
    conventional = batch / "batch_summary_grouped_fingerprints.xlsx"
    if conventional.is_file():
        return conventional
    recorded = manifest.get("interaction_fingerprints", {}).get("grouped_summary_xlsx")
    if recorded and native_path(recorded).is_file():
        return native_path(recorded).resolve()
    matches = list(batch.glob("*.xlsx"))
    if len(matches) != 1:
        raise ValueError("Expected exactly one workbook; supply --workbook for a renamed file")
    return matches[0]


def verify(batch, workbook=None, *, write_report=True):
    """Check saved values and layout; never repair or rewrite the workbook."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    batch = Path(batch).resolve()
    payload, manifest, failures = prepare(batch)
    path = workbook_path(batch, manifest, workbook)
    columns, rows = payload["columns"], payload["rows"]
    last_column, last_row = get_column_letter(len(columns)), 5 + len(rows)
    with closing(load_workbook(path, read_only=True, data_only=False)) as book:
        sheet = book["Fingerprints"]
        for r, actual in enumerate(sheet.iter_rows(min_row=3, max_row=last_row, max_col=len(columns)), 3):
            if r == 4:
                for c, cell in enumerate(actual):
                    starts_section = c == 0 or columns[c - 1]["section"] != columns[c]["section"]
                    expected = columns[c]["section"] if starts_section else None
                    assert cell.value == expected, f"Incorrect section label: {cell.coordinate}"
                continue
            source = payload.get("parent", {}) if r == 3 else None if r == 5 else rows[r - 6]
            for c, cell in enumerate(actual):
                expected = columns[c]["label"] if r == 5 else typed(source.get(columns[c]["key"]), columns[c])
                assert cell.data_type not in ("e", "f"), f"Unexpected error/formula at {cell.coordinate}"
                if isinstance(expected, (int, float)) and not isinstance(expected, bool):
                    assert type(cell.value) in (int, float) and math.isclose(cell.value, expected, rel_tol=1e-12, abs_tol=1e-10), cell.coordinate
                else:
                    assert cell.value == expected and (not isinstance(expected, bool) or isinstance(cell.value, bool)), cell.coordinate
        assert ("Failures" in book.sheetnames) == bool(failures)
        if failures:
            actual = list(book["Failures"].iter_rows(min_row=2, values_only=True))
            assert actual == [tuple(f[k] for k in ("candidate", "error", "expected", "observed", "diagnostic")) for f in failures]
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    auto_filter = root.find("m:autoFilter", ns)
    pane = root.find("m:sheetViews/m:sheetView/m:pane", ns)
    assert auto_filter is not None and auto_filter.attrib["ref"] == f"A5:{last_column}{last_row}", "Incorrect native filter range"
    assert pane is not None and all(pane.attrib.get(key) == value for key, value in {
        "xSplit": "2", "ySplit": "5", "topLeftCell": "C6", "state": "frozen",
    }.items()), "Incorrect freeze panes"
    merges = {item.attrib["ref"] for item in root.findall("m:mergeCells/m:mergeCell", ns)}
    groups = []
    for index, column in enumerate(columns, 1):
        if groups and groups[-1][0] == column["section"]:
            groups[-1][2] = index
        else:
            groups.append([column["section"], index, index])
    section_merges = {f"{get_column_letter(start)}4:{get_column_letter(end)}4"
                      for _, start, end in groups if end > start}
    assert section_merges <= merges, "Missing section band"
    result = {"status": "verified", "completed_count": int(manifest["completed_count"]), "failure_count": len(failures),
              "displayed_candidate_rows": len(rows),
              "candidate_count": manifest["candidate_count"], "columns": len(columns),
              "filter_range": f"A5:{last_column}{last_row}", "freeze_panes": "C6", "workbook": str(path)}
    if write_report:
        (batch / "xlsx_verification.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    return result


def export_batch_workbook(batch, output=None):
    """Export the saved payload, verify it, then record the successful path."""
    from .spreadsheet import export_workbook

    batch = Path(batch).resolve()
    payload, manifest, failures = prepare(batch)
    path = Path(output).resolve() if output is not None else batch / "batch_summary_grouped_fingerprints.xlsx"
    export_workbook(payload, path, failures=failures)
    verify(batch, path)
    manifest.setdefault("interaction_fingerprints", {})["grouped_summary_xlsx"] = str(path)
    manifest_path(batch).write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return path


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("batch", type=Path)
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--verify-only", action="store_true", help="Read an existing workbook without exporting")
    parser.add_argument("--workbook", type=Path, help="Existing workbook path for --verify-only")
    parser.add_argument("--output", type=Path, help="Export location; defaults to the batch folder")
    parser.add_argument("--sections", help="Comma-separated section names for a derived report")
    parser.add_argument("--columns", help="Comma-separated additional column keys for a derived report")
    parser.add_argument("--formats", help="Derived report formats: csv,xlsx, csv, or xlsx")
    parser.add_argument("--report-output", type=Path, help="Derived-report folder; defaults to BATCH/selected_report")
    args = parser.parse_args(argv)
    projected = any(value is not None for value in (args.sections, args.columns, args.formats, args.report_output))
    if projected and (args.prepare or args.verify_only or args.output):
        parser.error("Report selection cannot be combined with --prepare, --verify-only or --output")
    if args.workbook and not args.verify_only:
        parser.error("--workbook requires --verify-only")
    if args.prepare and (args.verify_only or args.output):
        parser.error("--prepare cannot be combined with export/verification options")
    if args.verify_only and args.output:
        parser.error("--verify-only cannot be combined with --output")
    if projected:
        from .report_options import export_selected_report
        print(f"Selected report: {export_selected_report(args.batch, args.sections, args.columns, args.formats or ('csv', 'xlsx'), args.report_output)}")
    elif args.prepare:
        payload, manifest, failures = prepare(args.batch)
        print(json.dumps({"failures": failures, "candidate_count": manifest["candidate_count"],
                          "completed_count": int(manifest["completed_count"]), "displayed_candidate_rows": len(payload["rows"])}))
    elif args.verify_only:
        print(json.dumps(verify(args.batch, args.workbook)))
    else:
        print(f"Verified grouped workbook: {export_batch_workbook(args.batch, args.output)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
