"""Optional Excel presentation of the grouped fingerprint summary."""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any


_COLORS = {
    "Identity": ("3C4856", "E9EDF1"),
    "Rosetta": ("3D617D", "E8F0F6"),
    "RosettaLigand": ("3D617D", "E8F0F6"),
    "Arpeggio": ("496955", "EAF1EB"),
    "PLIP": ("726144", "F3EEE6"),
    "Discovery Studio": ("6A5878", "EEEAF2"),
    "Structure & network": ("5D6670", "EBEEF0"),
    "Context": ("5D6670", "EBEEF0"),
    "HCDR3 pLDDT": ("385E6C", "E7F0F3"),
    "Boltz-2": ("4F6179", "ECF0F6"),
    "GNINA": ("796048", "F5EFE8"),
    "E2 architecture": ("6B6246", "F3EFE3"),
    "E3 discrimination": ("705F4B", "F3EEE8"),
    "Analysis status": ("5D6670", "EBEEF0"),
}
_TOOLS = {"Arpeggio", "PLIP", "Discovery Studio"}


def _sections(columns: list[dict[str, Any]], *, projected: bool = False) -> list[tuple[str, int, int]]:
    """Validate the schema and locate its contiguous, one-based bands."""
    if len(columns) < (3 if projected else 4) or any(not all(c.get(k) for k in ("key", "label", "section", "kind"))
                               for c in columns):
        raise ValueError("Columns require key, label, section and kind.")
    if len({c["key"] for c in columns}) != len(columns):
        raise ValueError("Column keys must be unique.")
    groups: list[tuple[str, int, int]] = []
    for index, column in enumerate(columns, 1):
        name = column["section"]
        if groups and groups[-1][0] == name:
            groups[-1] = (name, groups[-1][1], index)
        else:
            if any(group[0] == name for group in groups):
                raise ValueError(f"Section {name} must occupy contiguous columns.")
            groups.append((name, index, index))
    names = {name for name, _, _ in groups}
    if projected:
        if ({c["key"] for c in columns[:3]} != {"rank", "mutations", "candidate"}
                or any(c["section"] != "Identity" for c in columns[:3]) or names - _COLORS.keys()):
            raise ValueError("Projected reports require the three Identity columns and known sections.")
    elif (not _TOOLS <= names or len(names & {"Rosetta", "RosettaLigand"}) != 1
            or names - _COLORS.keys()):
        raise ValueError("Use the four tool sections and supported identity, context, confidence or score sections.")
    return groups


def _value(value: Any, kind: str, key: str) -> Any:
    if value is None or value in ("", "NA"):
        return None
    if kind == "text":
        return str(value)
    if kind == "boolean":
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.lower() in ("true", "false"):
            return value.lower() == "true"
        raise ValueError(f"Invalid boolean value for {key}: {value!r}")
    if kind not in ("integer", "number", "energy", "delta_energy") or isinstance(value, bool):
        raise ValueError(f"Invalid numeric kind or value for {key}: {value!r}")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid numeric value for {key}: {value!r}") from error
    if not math.isfinite(number) or (kind == "integer" and not number.is_integer()):
        raise ValueError(f"Invalid numeric value for {key}: {value!r}")
    return int(number) if kind == "integer" else number


def export_workbook(payload: dict[str, Any], output_path: str | Path, *,
                    failures: list[dict[str, Any]] | None = None) -> Path:
    """Write one candidate per row, with four tool bands and a pinned parent row.

    ``payload`` supplies title/note, columns (key/label/section/kind/width),
    parent values, candidate rows and an optional source_note. This exporter
    only formats those supplied values; it performs no scientific calculations.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError(
            "Excel export requires openpyxl. Install it with "
            "pip install 'fv-energy-analysis[reporting]'."
        ) from error

    columns, rows = payload["columns"], payload["rows"]
    groups = _sections(columns, projected=bool(payload.get("report_projection")))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fingerprints"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "34495E"
    last_column = len(columns)
    last_row = 5 + len(rows)

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3)
    sheet.cell(2, 1, str(payload.get("title") or "Interaction fingerprints"))
    sheet.cell(2, 1).font = Font(name="Arial", size=14, bold=True, color="263A4B")
    sheet.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    if last_column > 3:
        sheet.merge_cells(start_row=2, start_column=4, end_row=2, end_column=last_column)
        sheet.cell(2, 4, str(payload.get("note") or ""))
        sheet.cell(2, 4).font = Font(name="Arial", size=11, italic=True, color="596978")
        sheet.cell(2, 4).alignment = Alignment(vertical="center")
    for row_number, height in ((1, 10), (2, 46), (3, 30), (4, 27), (5, 62)):
        sheet.row_dimensions[row_number].height = height

    for name, start, end in groups:
        dark, light = _COLORS[name]
        if end > start:
            sheet.merge_cells(start_row=4, start_column=start, end_row=4, end_column=end)
        sheet.cell(4, start, name)
        for column_number in range(start, end + 1):
            for row_number, color in ((4, dark), (5, light)):
                cell = sheet.cell(row_number, column_number)
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(name="Arial", size=11, bold=True,
                                 color="FFFFFF" if row_number == 4 else "25313B")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    formats = {"integer": "0", "energy": "0.00;-0.00;0.00",
               "delta_energy": "+0.00;-0.00;0.00", "number": "0.000"}
    for column_number, column in enumerate(columns, 1):
        key, kind = column["key"], column["kind"]
        sheet.cell(5, column_number, column["label"])
        sheet.column_dimensions[get_column_letter(column_number)].width = column.get("width") or (
            34 if kind == "text" else 17 if kind == "boolean" else 15)
        for row_number, values in [(3, payload.get("parent") or {}), *enumerate(rows, 6)]:
            cell = sheet.cell(row_number, column_number, _value(values.get(key), kind, key))
            # Preserve filenames and any other supplied text literally, including '='.
            if kind == "text" and cell.value is not None:
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=11, bold=row_number == 3, color="25313B")
            alignment = "left" if kind == "text" else "center" if kind == "boolean" or key == "rank" else "right"
            cell.alignment = Alignment(horizontal=alignment, vertical="center", wrap_text=kind == "text")
            cell.number_format = formats.get(kind, "General")
            if row_number == 3 or row_number % 2:
                cell.fill = PatternFill("solid", fgColor="F0F2F4" if row_number == 3 else "F8FAFB")
    for row_number in range(6, last_row + 1):
        sheet.row_dimensions[row_number].height = 47
    sheet.freeze_panes = "C6"
    # A filter does not require unique labels, unlike an Excel Table.
    sheet.auto_filter.ref = f"A5:{get_column_letter(last_column)}{last_row}"

    if payload.get("source_note"):
        source_row = max(last_row + 3, 10)
        sheet.merge_cells(start_row=source_row, start_column=1, end_row=source_row,
                          end_column=min(8, last_column))
        cell = sheet.cell(source_row, 1, str(payload["source_note"]))
        cell.font = Font(name="Arial", size=11, italic=True, color="596978")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[source_row].height = 38

    if failures is None:
        failures = [{"candidate": row["candidate"], "error": row.get("error"),
                     "expected": None, "observed": None, "diagnostic": None}
                    for row in rows if row.get("status") == "failed"]
    if failures:
        failed = workbook.create_sheet("Failures")
        failed.sheet_view.showGridLines = False
        fields = ("candidate", "error", "expected", "observed", "diagnostic")
        labels = ("Complex", "Failure reason (PDB atom names)", "Expected InChIKey",
                  "Observed InChIKey", "Detailed diagnostic / source")
        failed.append(labels)
        failed.freeze_panes = "A2"
        failed.auto_filter.ref = f"A1:E{len(failures) + 1}"
        for column, width in enumerate((70, 85, 34, 34, 70), 1):
            failed.column_dimensions[get_column_letter(column)].width = width
            cell = failed.cell(1, column)
            cell.fill = PatternFill("solid", fgColor="6E4747")
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        failed.row_dimensions[1].height = 30
        for row_number, values in enumerate(failures, 2):
            failed.row_dimensions[row_number].height = 92
            for column, field in enumerate(fields, 1):
                value = values.get(field)
                cell = failed.cell(row_number, column, str(value) if value is not None else None)
                if value is not None:
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=11, color="25313B")
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
