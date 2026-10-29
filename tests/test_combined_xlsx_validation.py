"""Combined exports retain rejected rows and bind them to saved diagnostics."""
import csv
from hashlib import sha256
import importlib.util
import json
from pathlib import Path

import pytest

from fv_energy_analysis.confidence import confidence_columns
from fv_energy_analysis.fingerprint_reports import _grouped_columns
from fv_energy_analysis.model_scores import score_columns
from fv_energy_analysis.rescore_report import report_columns


from fv_energy_analysis import workbook_validation as verifier


requires_openpyxl = pytest.mark.skipif(
    importlib.util.find_spec("openpyxl") is None, reason="optional openpyxl is not installed")


def write_json(path, value):
    path.write_text(json.dumps(value), encoding="utf-8")


def save_table(root, payload, manifest):
    path = root / "batch_summary_grouped_fingerprints.csv"
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=[c["key"] for c in payload["columns"]])
        writer.writeheader()
        writer.writerows(payload["rows"])
    manifest.update(csv=str(path), csv_sha256=sha256(path.read_bytes()).hexdigest())
    write_json(root / "batch_summary_grouped_fingerprints.json", payload)
    write_json(root / "combined_manifest.json", manifest)


@pytest.fixture
def combined(tmp_path):
    root, source, diagnostic = tmp_path / "combined", tmp_path / "source", tmp_path / "diagnostic"
    for directory in (root, source, diagnostic):
        directory.mkdir()
    columns = report_columns(_grouped_columns() + confidence_columns() + score_columns()) + [
        {"key": "status", "label": "Analysis status", "section": "Analysis status", "kind": "text"},
        {"key": "error", "label": "Failure reason", "section": "Analysis status", "kind": "text"},
    ]
    empty = {column["key"]: None for column in columns}
    payload = {"columns": columns, "scoring_profile": "rosettaligand_legacy",
               "parent": {**empty, "candidate": "Parent", "status": "reference"},
               "rows": [{**empty, "rank": 1, "candidate": "valid.pdb", "status": "complete",
                         "rosetta_interface_REU": -15, "rosetta_interface_delta_REU": -1},
                        {**empty, "candidate": "rejected.pdb", "status": "failed", "error": "Mismatch"}]}
    manifest = {"status": "complete_with_failures", "candidate_count": 2,
                "completed_count": 1, "failure_count": 1}
    save_table(root, payload, manifest)
    with (root / "candidate_sources.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["candidate", "source", "source_rank"])
        writer.writeheader()
        writer.writerows({"candidate": row["candidate"], "source": str(source), "source_rank": row["rank"]}
                         for row in payload["rows"])
    write_json(source / "preparation_failures.json", [{"candidate": "rejected.pdb", "diagnostic_folder": str(diagnostic)}])
    write_json(diagnostic / "ligand_identity_diagnostic.json", {
        "differences": ["stereochemistry"], "expected": {"inchikey": "EXPECTED"},
        "observed": {"inchikey": "OBSERVED"},
        "stereocentre_differences": [{"atom_name": "C18", "expected_CIP": "S", "observed_CIP": "R"}],
    })
    return root, payload, manifest, diagnostic


def test_combined_preparation_keeps_failed_row_and_saved_diagnostics(combined):
    root, _, _, diagnostic = combined
    payload, manifest, failures = verifier.prepare(root)
    assert len(payload["rows"]) == manifest["candidate_count"] == 2
    assert manifest["completed_count"] == 1
    assert failures == [{"candidate": "rejected.pdb", "error": "Stereochemistry mismatch: C18: expected S, observed R",
                         "expected": "EXPECTED", "observed": "OBSERVED",
                         "diagnostic": str(diagnostic / "ligand_identity_diagnostic.json"),
                         "source": str(root.parent / "source")}]
    assert not (root / "batch_manifest.json").exists()


def test_user_csv_rename_is_resolved_by_hash_without_modifying_manifest(combined):
    root, _, _, _ = combined
    manifest = root / "combined_manifest.json"
    before = manifest.read_bytes()
    (root / "batch_summary_grouped_fingerprints.csv").rename(root / "renamed_by_user.csv")
    verifier.prepare(root)
    assert manifest.read_bytes() == before
    assert not (root / "batch_summary_grouped_fingerprints.csv").exists()


def test_changed_csv_is_rejected(combined):
    root, _, _, _ = combined
    path = root / "batch_summary_grouped_fingerprints.csv"
    path.write_text(path.read_text().replace("valid.pdb", "changed.pdb"))
    with pytest.raises(AssertionError, match="hash differs"):
        verifier.prepare(root)


def test_ambiguous_renamed_csv_is_rejected(combined):
    root, _, _, _ = combined
    path = root / "batch_summary_grouped_fingerprints.csv"
    data = path.read_bytes()
    path.unlink()
    (root / "one.csv").write_bytes(data)
    (root / "two.csv").write_bytes(data)
    with pytest.raises(AssertionError, match="exactly one"):
        verifier.prepare(root)


def test_csv_and_json_cell_mismatch_is_rejected(combined):
    root, payload, _, _ = combined
    payload["rows"][0]["rosetta_interface_REU"] = -16
    write_json(root / "batch_summary_grouped_fingerprints.json", payload)
    with pytest.raises(AssertionError, match="row differs"):
        verifier.prepare(root)


def test_rejected_row_cannot_contain_scientific_score(combined):
    root, payload, manifest, _ = combined
    payload["rows"][1]["rosetta_interface_REU"] = -20
    save_table(root, payload, manifest)
    with pytest.raises(AssertionError, match="Rejected candidate"):
        verifier.prepare(root)


def test_missing_provenance_candidate_is_rejected(combined):
    root, _, _, _ = combined
    (root / "candidate_sources.csv").write_text("candidate,source,source_rank\nvalid.pdb,source,1\n")
    with pytest.raises(AssertionError, match="provenance count"):
        verifier.prepare(root)


def test_optional_e3_section_keeps_original_scores_and_rejected_rows(combined):
    from fv_energy_analysis.e3_reports import append_e3_to_payload
    root, payload, manifest, _ = combined
    augmented = append_e3_to_payload(payload, {})
    save_table(root, augmented, manifest)
    result, _, failures = verifier.prepare(root)
    assert len(result["columns"]) > len(payload["columns"])
    assert result["rows"][0]["rosetta_interface_REU"] == -15
    assert result["rows"][0]["e3_status"] == "not_assessed"
    assert result["rows"][1]["e3_status"] == "excluded_input_rejection"
    assert len(failures) == 1


@requires_openpyxl
def test_full_current_schema_exports_and_verifies_with_failures(combined):
    from openpyxl import load_workbook
    from fv_energy_analysis.e3_reports import append_e3_to_payload

    root, payload, manifest, _ = combined
    augmented = append_e3_to_payload(payload, {})
    augmented["parent"]["e2_o3_hbond"] = False
    augmented["rows"][0]["e2_o3_hbond"] = True
    save_table(root, augmented, manifest)
    assert len(augmented["columns"]) == 124
    output = verifier.export_batch_workbook(root)
    report = verifier.verify(root, output, write_report=False)
    assert report["status"] == "verified"
    assert report["displayed_candidate_rows"] == 2
    assert report["columns"] == 124
    assert report["filter_range"] == "A5:DT7"
    workbook = load_workbook(output)
    sheet = workbook["Fingerprints"]
    positions = {column["key"]: i for i, column in enumerate(augmented["columns"], 1)}
    assert sheet.cell(4, 4).value == "RosettaLigand"
    assert sheet.cell(4, positions["e3_status"]).value == "E3 discrimination"
    assert sheet.cell(4, positions["status"]).value == "Analysis status"
    assert sheet.cell(3, positions["e2_o3_hbond"]).value is False
    assert sheet.cell(6, positions["e2_o3_hbond"]).value is True
    assert workbook["Failures"].cell(2, 1).value == "rejected.pdb"
    assert workbook["Failures"].cell(2, 2).value.startswith("Stereochemistry mismatch:")
    workbook.close()


@requires_openpyxl
def test_verification_does_not_repair_a_bad_filter(combined):
    from openpyxl import load_workbook

    root, _, _, _ = combined
    output = verifier.export_batch_workbook(root)
    workbook = load_workbook(output)
    workbook["Fingerprints"].auto_filter.ref = "A5:B7"
    workbook.save(output)
    workbook.close()
    original = output.read_bytes()
    with pytest.raises(AssertionError, match="filter range"):
        verifier.verify(root, output, write_report=False)
    assert output.read_bytes() == original


@requires_openpyxl
def test_verification_resolves_a_single_renamed_workbook(combined):
    root, _, _, _ = combined
    output = verifier.export_batch_workbook(root)
    renamed = output.with_name("user_named_final.xlsx")
    output.rename(renamed)
    manifest_before = (root / "combined_manifest.json").read_bytes()
    verifier.verify(root, write_report=False)
    assert (root / "combined_manifest.json").read_bytes() == manifest_before
    assert not output.exists()


@requires_openpyxl
def test_verification_checks_every_scientific_cell(combined):
    from openpyxl import load_workbook

    root, payload, _, _ = combined
    output = verifier.export_batch_workbook(root)
    index = next(i for i, c in enumerate(payload["columns"], 1) if c["key"] == "rosetta_interface_REU")
    workbook = load_workbook(output)
    workbook["Fingerprints"].cell(6, index, -999)
    workbook.save(output)
    workbook.close()
    with pytest.raises(AssertionError):
        verifier.verify(root, output, write_report=False)


@requires_openpyxl
def test_noncombined_report_can_omit_optional_metadata_sections(tmp_path):
    root = tmp_path / "ordinary"
    root.mkdir()
    columns = _grouped_columns(weights="ligand")
    payload = {"columns": columns, "scoring_profile": "rosettaligand_legacy",
               "parent": {"candidate": "Parent"},
               "rows": [{"candidate": "mutant.pdb", "rank": 1}]}
    write_json(root / "batch_summary_grouped_fingerprints.json", payload)
    write_json(root / "batch_manifest.json", {"candidate_count": 1, "completed_count": 1, "failure_count": 0})
    output = verifier.export_batch_workbook(root)
    assert verifier.verify(root, output, write_report=False)["columns"] == len(columns)
