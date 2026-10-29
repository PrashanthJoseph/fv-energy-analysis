"""Synthetic workbook checks; no scientific analysis or user reports are run."""

import importlib.util
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fv_energy_analysis.confidence import confidence_columns
from fv_energy_analysis.spreadsheet import export_workbook


def _payload():
    columns = [
        {"key": "rank", "label": "Rank", "section": "Identity", "kind": "integer", "width": 7},
        {"key": "candidate", "label": "Complex", "section": "Identity", "kind": "text", "width": 42},
        {"key": "mutations", "label": "Mutations", "section": "Identity", "kind": "text"},
        {"key": "energy", "label": "Interface", "section": "Rosetta", "kind": "energy"},
        {"key": "delta", "label": "Delta", "section": "Rosetta", "kind": "delta_energy"},
        {"key": "arpeggio_hb", "label": "H-bonds", "section": "Arpeggio", "kind": "integer"},
        {"key": "plip_hb", "label": "H-bonds", "section": "PLIP", "kind": "integer"},
        {"key": "ds_hb", "label": "H-bonds", "section": "Discovery Studio", "kind": "integer"},
        {"key": "warning", "label": "Warning", "section": "Context", "kind": "boolean"},
    ]
    return {"title": "Synthetic comparison", "note": "Counts by detector", "columns": columns,
            "parent": {"candidate": "reference.pdb", "energy": -12.25, "warning": False},
            "rows": [{"rank": "1", "candidate": "=literal_filename.pdb", "mutations": "E1F",
                      "energy": "-13.50", "delta": "-1.25", "arpeggio_hb": "2",
                      "plip_hb": 3, "ds_hb": "NA", "warning": "false"},
                     {"rank": 2, "candidate": "second.pdb", "energy": -12.0, "delta": 0.25,
                      "arpeggio_hb": 1, "plip_hb": 0, "ds_hb": 2, "warning": True}],
            "source_note": "Synthetic fixtures only"}


class MissingDependencyTests(unittest.TestCase):
    def test_optional_dependency_message(self):
        with patch.dict("sys.modules", {"openpyxl": None}):
            with self.assertRaisesRegex(RuntimeError, r"fv-energy-analysis\[reporting\]"):
                export_workbook(_payload(), Path("unused.xlsx"))


@unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional openpyxl is not installed")
class SpreadsheetTests(unittest.TestCase):
    def test_native_arpeggio_columns_and_missing_count_cells(self):
        from openpyxl import load_workbook
        from fv_energy_analysis.fingerprint_reports import _grouped_columns
        payload = {"columns": _grouped_columns(), "parent": {}, "rows": [
            {"candidate": "example", "arpeggio_ionic": 2, "arpeggio_CATIONPI": 1,
             "rosetta_hbond_total": None, "plip_hbond_total": 0}]}
        keys = [column["key"] for column in payload["columns"]]
        self.assertNotIn("arpeggio_salt_bridge", keys)
        self.assertNotIn("arpeggio_attractive_charge", keys)
        with TemporaryDirectory() as directory:
            output = export_workbook(payload, Path(directory) / "summary.xlsx")
            sheet = load_workbook(output)["Fingerprints"]
            col = lambda key: keys.index(key) + 1
            self.assertEqual(sheet.cell(5, col("arpeggio_ionic")).value, "ionic")
            self.assertEqual(sheet.cell(6, col("arpeggio_ionic")).value, 2)
            self.assertEqual(sheet.cell(6, col("arpeggio_CATIONPI")).value, 1)
            self.assertIsNone(sheet.cell(6, col("rosetta_hbond_total")).value)
            self.assertEqual(sheet.cell(6, col("plip_hbond_total")).value, 0)

    def test_architecture_section_roundtrip(self):
        from openpyxl import load_workbook
        from fv_energy_analysis.fingerprint_reports import ARCHITECTURE_COLUMNS
        payload = _payload()
        payload["columns"] += [{"key": key, "label": label, "section": "E2 architecture", "kind": kind}
                               for key, label, kind in ARCHITECTURE_COLUMNS]
        payload["rows"][0].update(e2_a_ring_pi_contact_count=2, e2_geometric_sandwich=True,
                                  e2_architecture_status="complete")
        payload["rows"][1].update(e2_architecture_status="partial: Discovery Studio output missing")
        with TemporaryDirectory() as directory:
            output = export_workbook(payload, Path(directory) / "summary.xlsx")
            sheet = load_workbook(output)["Fingerprints"]
            positions = {column["key"]: index for index, column in enumerate(payload["columns"], 1)}
            self.assertEqual(sheet.cell(4, 10).value, "E2 architecture")
            self.assertEqual(sheet.cell(6, positions["e2_a_ring_pi_contact_count"]).value, 2)
            self.assertIs(sheet.cell(6, positions["e2_geometric_sandwich"]).value, True)
            self.assertIsNone(sheet.cell(7, positions["e2_a_ring_pi_contact_count"]).value)
            self.assertTrue(sheet.auto_filter.ref.endswith("7"))

    def test_grouped_workbook_roundtrip(self):
        from openpyxl import load_workbook

        with TemporaryDirectory() as directory:
            output = Path(directory) / "nested" / "summary.xlsx"
            self.assertEqual(export_workbook(_payload(), output), output)
            book = load_workbook(output)
            sheet = book["Fingerprints"]
            self.assertEqual(book.sheetnames, ["Fingerprints"])
            self.assertEqual(sheet.freeze_panes, "C6")
            self.assertEqual(sheet.auto_filter.ref, "A5:I7")
            self.assertEqual(len(sheet.tables), 0)
            self.assertIn("D4:E4", {str(merged) for merged in sheet.merged_cells.ranges})
            self.assertEqual([sheet.cell(4, i).value for i in (4, 6, 7, 8)],
                             ["Rosetta", "Arpeggio", "PLIP", "Discovery Studio"])
            self.assertEqual(sheet["B3"].value, "reference.pdb")
            self.assertEqual(sheet["D3"].value, -12.25)
            self.assertEqual(sheet["B6"].value, "=literal_filename.pdb")
            self.assertEqual(sheet["B6"].data_type, "s")
            self.assertEqual(sheet["D6"].value, -13.5)
            self.assertEqual(sheet["E6"].value, -1.25)
            self.assertEqual(sheet["F6"].value, 2)
            self.assertEqual(sheet["F6"].data_type, "n")
            self.assertIsNone(sheet["H6"].value)
            self.assertIs(sheet["I6"].value, False)
            self.assertIs(sheet["I7"].value, True)
            self.assertEqual(sheet["A6"].alignment.horizontal, "center")
            self.assertEqual(sheet.column_dimensions["B"].width, 42)
            book.close()

    def test_confidence_section_is_numeric_and_fully_filterable(self):
        from openpyxl import load_workbook

        payload = _payload()
        confidence = confidence_columns()
        payload["columns"].extend(confidence)
        for record, start in ((payload["parent"], 82), (payload["rows"][0], 60),
                              (payload["rows"][1], 81)):
            record.update({"hcdr3_plddt_mean": str(start + 8), "hcdr3_plddt_min": str(start)})
            record.update({f"hcdr3_plddt_{number}": str(start + index)
                           for index, number in enumerate(range(97, 114))})
        with TemporaryDirectory() as directory:
            output = export_workbook(payload, Path(directory) / "confidence.xlsx")
            book = load_workbook(output)
            sheet = book["Fingerprints"]
            self.assertEqual(sheet["J4"].value, "HCDR3 pLDDT")
            self.assertIn("J4:AB4", {str(merged) for merged in sheet.merged_cells.ranges})
            self.assertEqual(sheet.auto_filter.ref, "A5:AB7")
            self.assertEqual(sheet.freeze_panes, "C6")
            self.assertEqual([sheet.cell(5, col).value for col in range(10, 29)],
                             [column["label"] for column in confidence])
            for row_number, expected in ((3, [90, 82, *range(82, 99)]),
                                         (6, [68, 60, *range(60, 77)]),
                                         (7, [89, 81, *range(81, 98)])):
                self.assertEqual([sheet.cell(row_number, col).value for col in range(10, 29)], expected)
                self.assertTrue(all(sheet.cell(row_number, col).data_type == "n" for col in range(10, 29)))
                self.assertTrue(all(sheet.cell(row_number, col).number_format == "0.000" for col in range(10, 29)))
            self.assertEqual([sheet.cell(4, col).value for col in (4, 6, 7, 8)],
                             ["Rosetta", "Arpeggio", "PLIP", "Discovery Studio"])
            self.assertEqual(sheet["D6"].value, -13.5)
            book.close()

    def test_invalid_schema_and_values(self):
        for case, expected in (("duplicate", "unique"), ("split", "contiguous"),
                               ("boolean", "boolean"), ("fraction", "numeric")):
            with self.subTest(case=case), TemporaryDirectory() as directory:
                payload = _payload()
                if case == "duplicate":
                    payload["columns"][1]["key"] = "rank"
                elif case == "split":
                    payload["columns"].append(dict(payload["columns"][3], key="other_energy"))
                elif case == "boolean":
                    payload["rows"][0]["warning"] = "invalid"
                else:
                    payload["rows"][0]["arpeggio_hb"] = 1.5
                with self.assertRaisesRegex(ValueError, expected):
                    export_workbook(payload, Path(directory) / "invalid.xlsx")

    def test_model_score_sections_extend_filter_and_keep_missing_parent_blank(self):
        from openpyxl import load_workbook
        from fv_energy_analysis.model_scores import score_columns

        payload = _payload()
        columns = score_columns()
        payload["columns"].extend(columns)
        for row in payload["rows"]:
            row.update({column["key"]: float(index) / 10 for index, column in enumerate(columns)})
        with TemporaryDirectory() as directory:
            output = export_workbook(payload, Path(directory) / "scores.xlsx")
            book = load_workbook(output)
            sheet = book["Fingerprints"]
            self.assertEqual(sheet["J4"].value, "Boltz-2")
            self.assertEqual(sheet["Y4"].value, "GNINA")
            self.assertEqual(sheet.auto_filter.ref, "A5:AH7")
            for column in range(10, 35):
                self.assertIsNone(sheet.cell(3, column).value)
                self.assertEqual(sheet.cell(6, column).data_type, "n")
                self.assertAlmostEqual(sheet.cell(6, column).value, (column - 10) / 10)
            self.assertEqual(sheet["D6"].value, -13.5)
            self.assertEqual(sheet.freeze_panes, "C6")
            book.close()


if __name__ == "__main__":
    unittest.main()
